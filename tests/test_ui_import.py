from openpyxl import Workbook
from sqlalchemy import select

from app.models import Group, Student
from app.ui.data import repository as repository_module
from app.io_tools.xlsx_config import XLSX_COLUMNS_BY_SHEET


def _workbook(path):
    workbook = Workbook()
    workbook.active.title = "Noise"
    workbook.active.append(["Random", "Unrelated"])
    workbook.active.append(["nothing", "useful"])
    return workbook


def test_smart_preview_imports_only_groups_and_students(tmp_path, monkeypatch):
    monkeypatch.setattr(repository_module, "DATABASE_URL", f"sqlite:///{tmp_path / 'smart.db'}")
    workbook = _workbook(tmp_path)
    groups = workbook.create_sheet("Roster")
    groups.append(["Группа", "Специальность"])
    groups.append(["IU7-11", "Computer science"])
    students = workbook.create_sheet("People")
    students.append(["Фамилия", "Имя", "Группа", "Почта"])
    students.append(["Ivanov", "Ivan", "IU7-11", "ivan@example.org"])
    marks = workbook.create_sheet("marks")
    marks.append(["student_id", "lesson_id", "data"])
    marks.append([1, 1, 80])
    source = tmp_path / "smart.xlsx"
    workbook.save(source)

    repo = repository_module.DiaryRepository()
    try:
        preview = repo.preview_import_from_xlsx(source, strategy="smart")
        assert preview.row_counts == {"groups": 1, "students": 1}
        assert repo.import_from_preview(preview) == {"groups": 1, "students": 1}
        student = repo.session.scalar(select(Student))
        assert student.bmstu_email == "ivan@example.org"
        assert repo.session.scalar(select(Group)).name == "IU7-11"
    finally:
        repo.close()


def test_template_preview_finds_model_table_and_skips_noise(tmp_path, monkeypatch):
    monkeypatch.setattr(repository_module, "DATABASE_URL", f"sqlite:///{tmp_path / 'template.db'}")
    workbook = _workbook(tmp_path)
    sheet = workbook.create_sheet("Somewhere")
    sheet["C5"] = "name"
    sheet["D5"] = "speciality"
    sheet["C6"] = "IU7-12"
    sheet["D6"] = "Mathematics"
    source = tmp_path / "template.xlsx"
    workbook.save(source)

    repo = repository_module.DiaryRepository()
    try:
        preview = repo.preview_import_from_xlsx(source, strategy="template")
        assert preview.row_counts == {"groups": 1}
        assert repo.import_from_preview(preview) == {"groups": 1}
        assert repo.session.scalar(select(Group)).name == "IU7-12"
    finally:
        repo.close()


def test_smart_merge_updates_existing_student_without_erasing_optional_fields(tmp_path, monkeypatch):
    monkeypatch.setattr(repository_module, "DATABASE_URL", f"sqlite:///{tmp_path / 'merge.db'}")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Roster"
    sheet.append(["Фамилия", "Имя", "Группа", "Почта"])
    sheet.append(["Ivanov", "Ivan", "IU7-11", "ivan@example.org"])
    source = tmp_path / "merge.xlsx"
    workbook.save(source)

    repo = repository_module.DiaryRepository()
    try:
        repo.import_from_preview(repo.preview_import_from_xlsx(source, strategy="smart"))
        student = repo.session.scalar(select(Student))
        student.patronymic = "Ivanovich"
        repo.session.commit()
        sheet["D2"] = "new@example.org"
        workbook.save(source)
        preview = repo.preview_import_from_xlsx(source, strategy="smart")
        assert preview.row_counts == {"students": 1}
        repo.import_from_preview(preview)
        assert len(repo.session.scalars(select(Student)).all()) == 1
        student = repo.session.scalar(select(Student))
        assert student.bmstu_email == "new@example.org"
        assert student.patronymic == "Ivanovich"
    finally:
        repo.close()


def test_standard_preview_rejects_invalid_data_before_import(tmp_path, monkeypatch):
    monkeypatch.setattr(repository_module, "DATABASE_URL", f"sqlite:///{tmp_path / 'strict.db'}")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "groups"
    sheet.append(XLSX_COLUMNS_BY_SHEET["groups"])
    sheet.append([None, None, None, None, None])
    sheet.append([None, "IU7-11", None, None, None])
    students = workbook.create_sheet("students")
    students.append(XLSX_COLUMNS_BY_SHEET["students"])
    students.append([None, "not-an-id", "Ivanov", "Ivan", None, None, None, None, None])
    source = tmp_path / "strict.xlsx"
    workbook.save(source)

    repo = repository_module.DiaryRepository()
    try:
        preview = repo.preview_import_from_xlsx(source, strategy="standard")
        assert preview.errors
        assert repo.session.scalars(select(Group)).all() == []
    finally:
        repo.close()


def test_template_accepts_missing_optional_email_and_skips_missing_group(tmp_path, monkeypatch):
    monkeypatch.setattr(repository_module, "DATABASE_URL", f"sqlite:///{tmp_path / 'optional.db'}")
    workbook = Workbook()
    groups = workbook.active
    groups.title = "Groups"
    groups.append(["name", "speciality"])
    groups.append(["IU7-11", "Math"])
    students = workbook.create_sheet("Students")
    students.append(["group_id", "surname", "first_name"])
    students.append([1, "Ivanov", "Ivan"])
    students.append([None, "Petrov", "Petr"])
    source = tmp_path / "optional.xlsx"
    workbook.save(source)

    repo = repository_module.DiaryRepository()
    try:
        preview = repo.preview_import_from_xlsx(source, strategy="template")
        assert preview.row_counts == {"groups": 1, "students": 1}
        assert preview.warnings
        repo.import_from_preview(preview)
        student = repo.session.scalar(select(Student))
        assert student.surname == "Ivanov"
        assert student.bmstu_email is None
    finally:
        repo.close()
